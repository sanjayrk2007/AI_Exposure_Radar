import os
import sys
import openpyxl

sys.path.insert(0, os.path.abspath("."))

try:
    from radar.contracts.models import VendorProfile, VendorVerdict
except ImportError:
    from contracts.models import VendorProfile, VendorVerdict
from m5_risk.criticality import assess_criticality
from m5_risk.planner import plan_depth
from m5_risk.risk import assess_risk
from m5_risk.render import render_depth_applied

TARGET_FILES = ["verify_no_llm.xlsx", "verify_with_llm.xlsx"]
col_map = {"L": 12, "M": 13, "N": 14, "S": 19, "T": 20, "U": 21, "V": 22}
ASSESSED_STAMP = "Team Meridian (M5 Risk Engine) / 2026-10-11"

def parse_coverage_sheet(ws_cov):
    """Parses the Coverage sheet using pure openpyxl (Zero pandas/numpy DLL dependencies)."""
    cov_dict = {}
    for row in ws_cov.iter_rows(min_row=2, values_only=True):
        vid = row[0]
        if not vid:
            continue
        if vid not in cov_dict:
            cov_dict[vid] = {
                "ai_usage": row[1],
                "strength": row[2],
                "coverage": float(row[3]) if row[3] is not None else 0.0,
                "sources": []
            }
        cov_dict[vid]["sources"].append({
            "source_type": row[4],
            "status": row[5],
            "pages_fetched": int(row[6]) if row[6] is not None else 0
        })
    return cov_dict

for target_name in TARGET_FILES:
    filepath = target_name
    if not os.path.exists(filepath):
        desktop_path = os.path.join(os.path.expanduser("~"), "Desktop", target_name)
        if os.path.exists(desktop_path):
            filepath = desktop_path
        else:
            print(f"[-] '{target_name}' not found. Skipping.")
            continue

    print(f"\n[+] Processing '{filepath}' (Pure Python mode)...")
    wb = openpyxl.load_workbook(filepath)
    ws = wb["Vendor Inventory"]
    ws_cov = wb["Coverage"]

    cov_data = parse_coverage_sheet(ws_cov)

    for r in range(6, 12):  # Rows 6 to 11 correspond to V-001 through V-006
        vendor = VendorProfile(
            vendor_id=ws.cell(row=r, column=2).value,
            name=ws.cell(row=r, column=3).value,
            description=ws.cell(row=r, column=4).value,
            service=ws.cell(row=r, column=5).value,
            category=ws.cell(row=r, column=6).value,
            website=ws.cell(row=r, column=7).value,
            process=ws.cell(row=r, column=8).value,
            dependency=ws.cell(row=r, column=9).value,
            data_class=ws.cell(row=r, column=10).value,
            data_volume=ws.cell(row=r, column=11).value
        )

        # 1. Criticality Tier & Rationale (Cols L & M)
        crit = assess_criticality(vendor)

        # 2. Depth Plan (Col N)
        depth = plan_depth(crit, vendor)
        cov_info = cov_data.get(vendor.vendor_id, {
            "ai_usage": "Inconclusive",
            "strength": "NONE",
            "coverage": 0.0,
            "sources": []
        })

        pages_sum = sum(s["pages_fetched"] for s in cov_info["sources"])
        statuses = "; ".join(f"{s['source_type']}:{s['status']}" for s in cov_info["sources"]) if cov_info["sources"] else "N/A"
        cov_ratio = cov_info["coverage"]

        plan_dict = {
            "max_pages": depth.max_pages,
            "max_depth": depth.max_depth,
            "sources": depth.sources,
            "stop_rule": depth.stop_rule
        }
        depth_applied = render_depth_applied(crit, pages_sum, ", ".join(depth.sources), statuses, cov_ratio, plan_dict)

        # 3. AI Security Risk Assessment (Cols S, T, U)
        strength_val = str(cov_info["strength"]) if cov_info["strength"] is not None else "NONE"
        ai_usage_val = str(cov_info["ai_usage"]) if cov_info["ai_usage"] is not None else "Inconclusive"

        verdict = VendorVerdict(
            vendor_id=vendor.vendor_id,
            ai_usage=ai_usage_val,
            strength=strength_val,
            coverage=cov_ratio,
            signal_ids=[],
            providers=[],
            use_cases=[],
            reasons=[]
        )
        risk = assess_risk(vendor, crit, verdict)

        # Write populated values into Excel cells
        ws.cell(row=r, column=col_map["L"]).value = crit.tier
        ws.cell(row=r, column=col_map["M"]).value = crit.rationale
        ws.cell(row=r, column=col_map["N"]).value = depth_applied
        ws.cell(row=r, column=col_map["S"]).value = risk.risk_class
        ws.cell(row=r, column=col_map["T"]).value = risk.rationale
        ws.cell(row=r, column=col_map["U"]).value = risk.recommended_action
        ws.cell(row=r, column=col_map["V"]).value = ASSESSED_STAMP

        print(f"  [+] {vendor.vendor_id} ({vendor.name}): Tier={crit.tier} | Risk={risk.risk_class}")

    out_file = filepath.replace(".xlsx", "_M5_completed.xlsx")
    wb.save(out_file)
    print(f"[✓] Saved completed deliverable: '{out_file}'")