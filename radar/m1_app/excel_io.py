"""M1 Excel input adapter: validate workbooks and build shared vendor profiles."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from radar.contracts.models import VendorProfile


REQUIRED_FIELDS = (
    "Vendor ID",
    "Vendor Name",
    "Vendor Description",
    "Service / Product Provided",
    "Category",
    "Public Website",
    "Business Process Supported",
    "Operational Dependency",
    "Data Classification Accessed",
    "Data Volume (annual)",
)
FIELD_MAP = {
    "vendor_id": "Vendor ID",
    "name": "Vendor Name",
    "description": "Vendor Description",
    "service": "Service / Product Provided",
    "category": "Category",
    "website": "Public Website",
    "process": "Business Process Supported",
    "dependency": "Operational Dependency",
    "data_class": "Data Classification Accessed",
    "data_volume": "Data Volume (annual)",
}
EXAMPLE_VENDOR_ID = "V-000"


def _clean(value: Any) -> str:
    """Normalize Excel values into strings expected by the shared contract."""
    return "" if value is None else str(value).strip()


def read_vendor_profiles(
    path: str | Path,
    sheet_name: str = "Vendor Inventory",
) -> list[VendorProfile]:
    """Read vendor rows without modifying the workbook.

    Input columns B:K and header row 4 follow radar.input_inspection.
    The example vendor V-000 is excluded.
    """
    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    workbook = load_workbook(
        filename=workbook_path,
        read_only=True,
        data_only=True,
    )

    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Required worksheet not found: {sheet_name!r}")

        sheet = workbook[sheet_name]
        headers = [
            _clean(sheet.cell(row=4, column=column).value)
            for column in range(2, 12)
        ]

        if tuple(headers) != REQUIRED_FIELDS:
            raise ValueError(
                "Unexpected headers in B4:K4. "
                f"Expected {REQUIRED_FIELDS!r}, got {tuple(headers)!r}"
            )

        profiles: list[VendorProfile] = []
        seen_ids: set[str] = set()

        for row_number in range(5, sheet.max_row + 1):
            values = {
                headers[index]: _clean(
                    sheet.cell(row=row_number, column=index + 2).value
                )
                for index in range(len(headers))
            }

            if not any(values.values()):
                continue

            vendor_id = values["Vendor ID"]
            if vendor_id == EXAMPLE_VENDOR_ID:
                continue
            if not vendor_id:
                raise ValueError(
                    f"Missing Vendor ID in row {row_number}"
                )
            if vendor_id in seen_ids:
                raise ValueError(
                    f"Duplicate Vendor ID {vendor_id!r} in row {row_number}"
                )

            if not values["Vendor Name"]:
                raise ValueError(
                    f"Missing Vendor Name for {vendor_id} in row {row_number}"
                )

            seen_ids.add(vendor_id)
            fields = {
                attribute: values[column]
                for attribute, column in FIELD_MAP.items()
            }
            profiles.append(VendorProfile(**fields))

        return profiles
    finally:
        workbook.close()
