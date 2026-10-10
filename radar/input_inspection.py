"""Read vendor input rows without modifying the workbook."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

INPUT_COLUMNS = tuple(range(2, 12))  # B:K, inclusive
EXAMPLE_VENDOR_ID = "V-000"


def read_vendor_profiles(path: str | Path, sheet_name: str = "Vendor Inventory") -> list[dict[str, Any]]:
    """Read B:K vendor data rows, excluding the worked-example V-000 row.

    The source workbook is opened read-only and is never saved or modified.
    Header names are taken from row 4; records start on row 5.
    """
    workbook = load_workbook(filename=Path(path), read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Required worksheet not found: {sheet_name!r}")
        sheet = workbook[sheet_name]
        headers = [sheet.cell(row=4, column=column).value for column in INPUT_COLUMNS]
        if not all(isinstance(header, str) and header.strip() for header in headers):
            raise ValueError("Expected non-empty headers in B4:K4")

        records: list[dict[str, Any]] = []
        for row_number in range(5, sheet.max_row + 1):
            values = [sheet.cell(row=row_number, column=column).value for column in INPUT_COLUMNS]
            if not any(value is not None for value in values):
                continue
            record = dict(zip(headers, values))
            if record.get("Vendor ID") == EXAMPLE_VENDOR_ID:
                continue
            records.append(record)
        return records
    finally:
        workbook.close()
