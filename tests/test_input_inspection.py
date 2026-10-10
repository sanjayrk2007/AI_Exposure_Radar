from openpyxl import Workbook

from radar.input_inspection import read_vendor_profiles


def make_workbook(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vendor Inventory"
    headers = [
        "Vendor ID", "Vendor Name", "Vendor Description",
        "Service / Product Provided", "Category", "Public Website",
        "Business Process Supported", "Operational Dependency",
        "Data Classification Accessed", "Data Volume (annual)",
    ]
    for index, header in enumerate(headers, start=2):
        sheet.cell(row=4, column=index, value=header)
    sheet.cell(row=5, column=2, value="V-000")
    sheet.cell(row=5, column=3, value="Worked example")
    sheet.cell(row=6, column=2, value="V-001")
    sheet.cell(row=6, column=3, value="Test Vendor")
    workbook.save(path)


def test_excludes_worked_example_and_reads_b_to_k(tmp_path):
    path = tmp_path / "input.xlsx"
    make_workbook(path)

    records = read_vendor_profiles(path)

    assert len(records) == 1
    assert records[0]["Vendor ID"] == "V-001"
    assert records[0]["Vendor Name"] == "Test Vendor"
    assert "Criticality Tier" not in records[0]


def test_missing_worksheet_fails_clearly(tmp_path):
    path = tmp_path / "input.xlsx"
    make_workbook(path)

    try:
        read_vendor_profiles(path, sheet_name="Missing")
    except ValueError as error:
        assert "Missing" in str(error)
    else:
        raise AssertionError("Expected ValueError for missing worksheet")
