import hashlib
import json

import pytest
import requests
from openpyxl import Workbook, load_workbook

from m4_intel import run_demo
from m4_intel.adjudicate import LocalLLM, require_local_url, validate_proposal
from m4_intel.analyse import (
    DEFAULT_PLANNED_SOURCES,
    analyse,
    analyse_detailed,
    check_document_integrity,
)
from radar.collector import (
    RobotsPolicy,
    build_source_reports,
    classify_source_type,
    collect_vendor_pages,
)
from radar.contracts.models import Document, SourceReport
from radar.export import export_workbook

TECH = "Our fraud engine uses an OpenAI model via the API to score each transaction."
SITE = "www.vendor-one.test"
HEADERS = [
    "Vendor ID", "Vendor Name", "Vendor Description", "Service", "Category", "Public Website",
    "Process", "Dependency", "Data", "Volume", "Criticality Tier", "Criticality Rationale",
    "Depth", "AI Usage Detected", "Evidence", "How", "Providers", "Risk", "Risk Rationale",
    "Action", "Assessed By",
]


def make_doc(doc_id, source_type, text, url=None, vendor_id="V-001", content_hash=None):
    return {
        "doc_id": doc_id, "vendor_id": vendor_id, "source_type": source_type,
        "url": url or f"https://{SITE}/{doc_id}", "fetched_at": "2026-10-11T08:00:00+00:00",
        "http_status": 200, "content_hash": content_hash or hashlib.sha256(text.encode()).hexdigest(),
        "title": "t", "text": text, "snapshot_path": "", "meta": {},
    }


def write_input_workbook(path, vendors=("V-001", "V-002")):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vendor Inventory"
    sheet["B2"] = "Invented test inventory"
    for offset, header in enumerate(HEADERS):
        sheet.cell(row=4, column=2 + offset, value=header)
    sheet.cell(row=5, column=2, value="V-000")
    sheet.cell(row=5, column=3, value="EXAMPLE")
    for column in range(12, 23):
        sheet.cell(row=5, column=column, value=f"example-{column}")
    for index, vendor_id in enumerate(vendors):
        row = 6 + index
        sheet.cell(row=row, column=2, value=vendor_id)
        sheet.cell(row=row, column=3, value=f"Vendor {index + 1}")
        sheet.cell(row=row, column=7, value=SITE if index == 0 else "www.vendor-two.test")
        sheet.cell(row=row, column=9, value="High")
    workbook.save(path)
    return path


@pytest.fixture(autouse=True)
def recorded_sleep(monkeypatch):
    sleeps = []
    monkeypatch.setattr("radar.collector.time.sleep", sleeps.append)
    return sleeps


class FakeResponse:
    def __init__(self, url, text="", status=200, content_type="text/html"):
        self.url, self.text, self.status_code = url, text, status
        self.headers = {"Content-Type": content_type}


class FakeSession:
    def __init__(self, responses):
        self.responses, self.headers, self.requested = responses, {}, []

    def get(self, url, **kwargs):
        self.requested.append(url)
        if url in self.responses:
            response = self.responses[url]
            if isinstance(response, Exception):
                raise response
            return response
        return FakeResponse(url, "", status=404)

    def close(self):
        pass


def profile():
    return {"Vendor ID": "V-001", "Vendor Name": "Vendor One", "Public Website": SITE}


def test_web_sentence_naming_provider_never_yields_no():
    docs = [make_doc("D1", "WEB", TECH)]
    reports = [SourceReport("V-001", s, "OK", 1, "") for s in ("TRUST", "DOCS", "WEB")]
    docs += [make_doc("D2", "TRUST", "Plain unrelated text."), make_doc("D3", "DOCS", "More unrelated text.")]
    signals, verdict = analyse({"vendor_id": "V-001"}, docs, reports, ["TRUST", "DOCS", "WEB"], use_model=False)
    assert signals and verdict.coverage == 1.0
    assert verdict.ai_usage == "Inconclusive"


def test_default_plan_does_not_inflate_coverage_when_only_web_collected():
    _, verdict = analyse({"vendor_id": "V-001"}, [make_doc("D1", "WEB", TECH)], [], DEFAULT_PLANNED_SOURCES, use_model=False)
    assert verdict.coverage == pytest.approx(1.5 / 10.5)
    assert verdict.ai_usage == "Inconclusive"


def test_integrity_rejects_bad_hash_foreign_host_and_bad_url():
    good = make_doc("D1", "TRUST", TECH)
    assert check_document_integrity(good, "vendor-one.test") == []
    assert check_document_integrity(make_doc("D2", "TRUST", TECH, content_hash="abc"), "vendor-one.test")
    assert check_document_integrity(make_doc("D3", "TRUST", TECH, url="https://evil.test/x"), "vendor-one.test")
    assert check_document_integrity(make_doc("D4", "TRUST", TECH, url="ftp://www.vendor-one.test/x"), "vendor-one.test")
    assert check_document_integrity(make_doc("D5", "TRUST", TECH, url="https://docs.vendor-one.test/x"), "vendor-one.test") == []


def test_strict_analysis_excludes_failing_documents_and_says_why():
    docs = [make_doc("D1", "TRUST", TECH), make_doc("D2", "DOCS", TECH, url="https://other.test/p")]
    result = analyse_detailed({"vendor_id": "V-001", "Public Website": SITE}, docs, [], ["TRUST", "DOCS"],
                              use_model=False, strict_integrity=True)
    assert set(result.rejected_documents) == {"D2"}
    assert {s.doc_id for s in result.signals} == {"D1"}
    assert any("integrity" in reason for reason in result.verdict.reasons)


def test_validate_proposal_requires_exact_quote_and_filters_providers():
    good = {"is_technical_evidence": True, "reason_code": "NAMES_PROVIDER", "providers": ["OpenAI", "Invented Corp"],
            "quote": "uses an OpenAI model"}
    proposal = validate_proposal(good, TECH)
    assert proposal["providers"] == ["OpenAI"] and proposal["review_status"] == "LLM_PROPOSED_UNREVIEWED"
    assert validate_proposal({**good, "quote": "uses a Mistral model"}, TECH) is None
    assert validate_proposal({**good, "quote": ""}, TECH) is None
    assert validate_proposal({**good, "reason_code": "MADE_UP"}, TECH) is None
    assert validate_proposal({**good, "is_technical_evidence": "yes"}, TECH) is None
    assert validate_proposal("not json", TECH) is None


@pytest.mark.parametrize("url", ["http://example.com:11434", "https://api.openai.com", "http://192.168.1.5:11434",
                                 "http://user:pw@127.0.0.1:11434"])
def test_non_local_llm_urls_are_refused(url):
    with pytest.raises(ValueError):
        require_local_url(url)
    with pytest.raises(ValueError):
        LocalLLM(base_url=url)


class FakeLLMSession:
    def __init__(self, content=None, error=None, status=200):
        self.content, self.error, self.status, self.posts = content, error, status, []

    def post(self, url, json=None, timeout=None):
        self.posts.append((url, json))
        if self.error:
            raise self.error

        class Response:
            status_code = self.status

            def json(inner):
                return {"message": {"content": self.content}}
        return Response()


def good_content(quote="uses an OpenAI model"):
    return json.dumps({"is_technical_evidence": True, "reason_code": "NAMES_PROVIDER",
                       "providers": ["OpenAI"], "quote": quote})


def test_local_llm_caches_and_is_deterministic(tmp_path):
    session = FakeLLMSession(good_content())
    llm = LocalLLM("http://127.0.0.1:11434", "test-model", tmp_path, session=session)
    first = llm.propose(TECH)
    second = llm.propose(TECH)
    assert first == second and len(session.posts) == 1 and llm.stats["cache_hits"] == 1
    assert session.posts[0][1]["options"]["temperature"] == 0
    assert session.posts[0][0].startswith("http://127.0.0.1")


def test_local_llm_unavailable_falls_back_without_repeated_calls(tmp_path):
    session = FakeLLMSession(error=requests.ConnectionError("down"))
    llm = LocalLLM("http://localhost:11434", "m", tmp_path, session=session)
    assert llm.propose("a sentence about AI") is None
    assert llm.propose("another sentence about AI") is None
    assert len(session.posts) == 1 and llm.available is False


def test_llm_quote_not_in_sentence_is_rejected(tmp_path):
    llm = LocalLLM("http://127.0.0.1:11434", "m", tmp_path, session=FakeLLMSession(good_content("uses Claude")))
    assert llm.propose(TECH) is None and llm.stats["invalid"] == 1


def test_adjudicator_can_raise_but_not_lower_specificity():
    docs = [make_doc("D1", "TRUST", "AI is used to help with documents.")]
    base_signals, _ = analyse({"vendor_id": "V-001"}, docs, [], ["TRUST"], use_model=False)
    base = base_signals[0].specificity
    sentence = base_signals[0].excerpt

    def raising(_):
        return {"is_technical_evidence": True, "reason_code": "CURRENT_SPECIFIC_USE", "providers": [],
                "quote": "AI is used", "review_status": "LLM_PROPOSED_UNREVIEWED"}

    def lowering(_):
        return {"is_technical_evidence": False, "reason_code": "VAGUE_MARKETING", "providers": [],
                "quote": "", "review_status": "LLM_PROPOSED_UNREVIEWED"}

    up = analyse_detailed({"vendor_id": "V-001"}, docs, [], ["TRUST"], use_model=False, adjudicator=raising)
    down = analyse_detailed({"vendor_id": "V-001"}, docs, [], ["TRUST"], use_model=False, adjudicator=lowering)
    assert base < 0.7 and sentence
    if up.ambiguous:
        assert up.signals[0].specificity >= base and up.signals[0].method == "LLM"
        assert all(p["review_status"] == "LLM_PROPOSED_UNREVIEWED" for p in up.llm_proposed.values())
    assert down.signals[0].specificity == base


def test_classify_source_type_uses_path_only():
    assert classify_source_type("https://x.test/trust") == "TRUST"
    assert classify_source_type("https://x.test/legal/subprocessors") == "TRUST"
    assert classify_source_type("https://x.test/developers/api") == "DOCS"
    assert classify_source_type("https://x.test/ai-platform") == "WEB"
    assert classify_source_type("https://x.test/") == "WEB"


def test_robots_disallow_prevents_request_and_is_reported_blocked(tmp_path):
    home = "https://www.vendor-one.test/"
    session = FakeSession({
        "https://www.vendor-one.test/robots.txt": FakeResponse("", "User-agent: *\nDisallow: /trust\n", content_type="text/plain"),
        home: FakeResponse(home, "<html><body><p>Welcome.</p></body></html>"),
    })
    docs, events = collect_vendor_pages(profile(), tmp_path, session=session, probe_paths=True)
    assert "https://www.vendor-one.test/trust" not in session.requested
    assert any(e["status"] == "ROBOTS_DISALLOWED" for e in events)
    reports = {r.source_type: r for r in build_source_reports("V-001", docs, events)}
    assert reports["WEB"].status == "OK"
    assert reports["TRUST"].status == "BLOCKED"
    assert reports["DOCS"].status == "EMPTY" and "not evidence of no AI" in reports["DOCS"].note


def test_unreachable_robots_blocks_all_collection(tmp_path):
    session = FakeSession({"https://www.vendor-one.test/robots.txt": FakeResponse("", status=503)})
    docs, events = collect_vendor_pages(profile(), tmp_path, session=session)
    assert docs == [] and events[0]["status"] == "ROBOTS_UNREACHABLE"
    assert session.requested == ["https://www.vendor-one.test/robots.txt"]
    assert all(r.status == "BLOCKED" for r in build_source_reports("V-001", docs, events))


def test_robots_404_means_no_restrictions_and_exception_blocks():
    open_policy = RobotsPolicy("https://a.test/", FakeSession({}), 1)
    assert open_policy.allowed("https://a.test/anything")
    broken = RobotsPolicy("https://a.test/", FakeSession({"https://a.test/robots.txt": requests.ConnectionError("x")}), 1)
    assert not broken.allowed("https://a.test/anything")


def test_probe_paths_collect_typed_documents(tmp_path):
    home, trust, docs_url = ("https://www.vendor-one.test/", "https://www.vendor-one.test/trust",
                             "https://www.vendor-one.test/developers")
    session = FakeSession({
        home: FakeResponse(home, "<html><body><p>Home.</p></body></html>"),
        trust: FakeResponse(trust, f"<html><body><p>{TECH}</p></body></html>"),
        docs_url: FakeResponse(docs_url, "<html><body><p>Developer reference.</p></body></html>"),
    })
    docs, events = collect_vendor_pages(profile(), tmp_path, max_pages=6, session=session, probe_paths=True)
    assert {d.source_type for d in docs} == {"WEB", "TRUST", "DOCS"}
    assert {r.source_type: r.status for r in build_source_reports("V-001", docs, events)} == {
        "TRUST": "OK", "DOCS": "OK", "WEB": "OK"}


def test_export_refuses_same_path_and_preserves_provided_cells(tmp_path):
    source = write_input_workbook(tmp_path / "input.xlsx")
    from m4_intel.analyse import analyse_detailed as run_analysis
    from radar.export import VendorResult
    docs = [make_doc("D1", "TRUST", TECH)]
    reports = [SourceReport("V-001", "TRUST", "OK", 1, "ok")]
    outcome = run_analysis({"vendor_id": "V-001", "Public Website": SITE}, docs, reports, ["TRUST"],
                           use_model=False, strict_integrity=True)
    results = {"V-001": VendorResult("V-001", outcome.verdict, outcome.signals, reports, docs)}
    with pytest.raises(ValueError):
        export_workbook(source, source, results)
    original = source.read_bytes()
    output = export_workbook(source, tmp_path / "out.xlsx", results)
    assert source.read_bytes() == original
    out_book, in_book = load_workbook(output), load_workbook(source)
    out_sheet, in_sheet = out_book["Vendor Inventory"], in_book["Vendor Inventory"]
    for column in range(12, 23):
        assert out_sheet.cell(row=5, column=column).value == in_sheet.cell(row=5, column=column).value
    assert out_sheet["O6"].value == "Yes"
    assert TECH in out_sheet["P6"].value and "retrieved 2026-10-11" in out_sheet["P6"].value
    assert out_sheet["O7"].value is None
    # M5-dependent fields remain blank when no M5 results are supplied.
    for column in (12, 13, 14, 19, 20, 21):
        assert out_sheet.cell(row=6, column=column).value is None

    # Assessment metadata is populated by the exporter.
    assert "AI Exposure Radar team" in out_sheet["V6"].value
    assert "/" in out_sheet["V6"].value
    log = out_book["Evidence Log"]
    rows = list(log.iter_rows(min_row=2, values_only=True))
    assert rows and all(row[19] == "RULE_PROPOSED_UNREVIEWED" and row[20] in (None, "") for row in rows)
    assert all(row[0] != "V-000" for row in rows)



def test_export_populates_m5_fields_and_preserves_vendor_inputs(tmp_path):
    from radar.contracts.models import (
        CriticalityResult, DepthPlan, RiskAssessment, VendorProfile,
        VendorVerdict,
    )
    from radar.export import VendorResult

    source = write_input_workbook(tmp_path / "input.xlsx")
    original = load_workbook(source)
    original_sheet = original["Vendor Inventory"]
    original_vendor_cells = [
        original_sheet.cell(row=6, column=c).value for c in range(2, 12)
    ]
    original_example = [
        original_sheet.cell(row=5, column=c).value for c in range(1, 23)
    ]

    docs = [make_doc("D1", "TRUST", TECH)]
    reports = [SourceReport("V-001", "TRUST", "OK", 2, "ok")]
    outcome = analyse_detailed(
        {"vendor_id": "V-001", "Public Website": SITE},
        docs,
        reports,
        ["TRUST"],
        use_model=False,
        strict_integrity=True,
    )

    profile = VendorProfile(
        vendor_id="V-001",
        name="Vendor 1",
        description="Fraud detection service",
        service="Transaction fraud detection",
        category="Technology",
        website=SITE,
        process="Payment processing",
        dependency="High",
        data_class="Confidential",
        data_volume="High",
    )
    result = VendorResult(
        vendor_id="V-001",
        verdict=VendorVerdict(
            vendor_id="V-001",
            ai_usage=outcome.verdict.ai_usage,
            strength=outcome.verdict.strength,
            coverage=outcome.verdict.coverage,
            signal_ids=outcome.verdict.signal_ids,
            providers=outcome.verdict.providers,
            use_cases=["fraud detection"],
            reasons=outcome.verdict.reasons,
        ),
        signals=outcome.signals,
        reports=reports,
        documents=docs,
        criticality=CriticalityResult(
            vendor_id="V-001",
            tier="Critical",
            scores={"D": 4, "S": 5, "P": 3, "R": 4, "V": 5},
            floors_applied=["Critical process floor"],
            rationale="Supports a critical payment process.",
        ),
        depth_plan=DepthPlan(
            vendor_id="V-001",
            tier="Critical",
            sources=["TRUST", "DOCS"],
            max_pages=20,
            max_depth=3,
            stop_rule="Stop when sufficient evidence is collected.",
        ),
        risk=RiskAssessment(
            vendor_id="V-001",
            risk_class="High",
            rationale="AI processes sensitive payment data.",
            recommended_action="Review controls and obtain assurance evidence.",
        ),
        vendor_profile=profile,
        assessed_by="Integration Test",
        assessed_at="2026-10-13T00:00:00+00:00",
    )

    output = export_workbook(
        source, tmp_path / "results.xlsx", {"V-001": result}
    )
    workbook = load_workbook(output)
    sheet = workbook["Vendor Inventory"]

    assert sheet["L6"].value == "Critical"
    assert "Supports a critical payment process" in sheet["M6"].value
    assert sheet["N6"].value and "Critical" in sheet["N6"].value
    assert sheet["O6"].value == outcome.verdict.ai_usage
    assert TECH in sheet["P6"].value
    assert "Transaction fraud detection" in sheet["Q6"].value
    assert "Payment processing" in sheet["Q6"].value
    assert sheet["R6"].value is not None
    assert sheet["S6"].value == "High"
    assert "sensitive payment data" in sheet["T6"].value
    assert "obtain assurance evidence" in sheet["U6"].value
    assert sheet["V6"].value == (
        "Integration Test / 2026-10-13T00:00:00+00:00"
    )

    # Original vendor inputs and the example row must remain unchanged.
    assert [
        sheet.cell(row=6, column=c).value for c in range(2, 12)
    ] == original_vendor_cells
    assert [
        sheet.cell(row=5, column=c).value for c in range(1, 23)
    ] == original_example


def test_run_demo_cache_mode_isolates_failures_and_is_reproducible(tmp_path):
    source = write_input_workbook(tmp_path / "input.xlsx")
    cache = tmp_path / "cache"
    docs = [make_doc("D1", "TRUST", TECH), make_doc("D2", "DOCS", "Plain text."), make_doc("D3", "WEB", "Plain text.")]
    reports = [SourceReport("V-001", s, "OK", 1, "") for s in ("TRUST", "DOCS", "WEB")]
    target = cache / "V-001" / "collected.json"
    target.parent.mkdir(parents=True)
    target.write_text(json.dumps({"documents": docs, "reports": [r.__dict__ for r in reports], "events": []}))

    first = run_demo.run(source, tmp_path / "a.xlsx", cache_dir=cache, mode="cache", use_llm=False)
    second = run_demo.run(source, tmp_path / "b.xlsx", cache_dir=cache, mode="cache", use_llm=False)
    by_id = {line["vendor_id"]: line for line in first["vendors"]}
    assert by_id["V-001"]["ai_usage"] == "Yes" and by_id["V-001"]["error"] == ""
    assert by_id["V-002"]["ai_usage"] == "Inconclusive" and "No cached collection" in by_id["V-002"]["error"]
    assert "V-000" not in by_id

    def values(path):
        sheet = load_workbook(path)["Vendor Inventory"]
        return [[sheet.cell(row=r, column=c).value for c in range(15, 19)] for r in (6, 7)]

    assert values(tmp_path / "a.xlsx") == values(tmp_path / "b.xlsx")
    failed_row = values(tmp_path / "a.xlsx")[1]
    assert failed_row[0] == "Inconclusive" and "failed" in failed_row[1].lower()


def test_run_demo_rejects_low_delay_and_same_output(tmp_path):
    source = write_input_workbook(tmp_path / "input.xlsx")
    with pytest.raises(ValueError):
        run_demo.run(source, tmp_path / "o.xlsx", cache_dir=tmp_path / "c", delay=0.5, use_llm=False)
    with pytest.raises(ValueError):
        run_demo.run(source, source, cache_dir=tmp_path / "c", mode="cache", use_llm=False)


def test_collect_vendor_pages_enforces_minimum_delay_and_waits(tmp_path, recorded_sleep):
    with pytest.raises(ValueError):
        collect_vendor_pages(profile(), tmp_path, delay_seconds=0.5, session=FakeSession({}))
    with pytest.raises(ValueError):
        collect_vendor_pages(profile(), tmp_path, delay_seconds=0, session=FakeSession({}))
    home, trust = "https://www.vendor-one.test/", "https://www.vendor-one.test/trust"
    session = FakeSession({
        home: FakeResponse(home, "<html><body><p>Home.</p></body></html>"),
        trust: FakeResponse(trust, "<html><body><p>Trust page.</p></body></html>"),
    })
    collect_vendor_pages(profile(), tmp_path, max_pages=2, session=session, probe_paths=True)
    assert recorded_sleep and all(delay > 1.0 for delay in recorded_sleep)


def test_robots_crawl_delay_can_increase_but_never_reduce_the_wait(tmp_path, recorded_sleep):
    home = "https://www.vendor-one.test/"
    session = FakeSession({
        "https://www.vendor-one.test/robots.txt": FakeResponse("", "User-agent: *\nCrawl-delay: 5\n", content_type="text/plain"),
        home: FakeResponse(home, "<html><body><p>Home.</p></body></html>"),
        "https://www.vendor-one.test/trust": FakeResponse("https://www.vendor-one.test/trust", "<html><body><p>T.</p></body></html>"),
    })
    collect_vendor_pages(profile(), tmp_path, max_pages=2, session=session, probe_paths=True)
    assert recorded_sleep and max(recorded_sleep) > 4.0


def test_review_status_follows_proposer_and_is_never_human_reviewed():
    from radar.export import review_status
    base = dict(signal_id="S1", vendor_id="V-001", doc_id="D", source_type="TRUST", url="https://x.test/",
                excerpt="e", char_start=0, char_end=1, category="c", providers=[], specificity=0.5,
                weight=1.0, score=0.5, tag="STRONG", verified=True)
    from radar.contracts.models import Signal
    assert review_status(Signal(method="RULE", **base), {}) == "RULE_PROPOSED_UNREVIEWED"
    assert review_status(Signal(method="ML", **base), {}) == "ML_PROPOSED_UNREVIEWED"
    assert review_status(Signal(method="LLM", **base), {}) == "LLM_PROPOSED_UNREVIEWED"
    assert review_status(Signal(method="RULE", **base), {"S1": {"reason_code": "VAGUE_MARKETING"}}) == "LLM_PROPOSED_UNREVIEWED"
    assert all("HUMAN" not in review_status(Signal(method=m, **base), {}) for m in ("RULE", "ML", "LLM", "MANUAL"))


def test_model_and_provider_names_are_not_assumed_to_be_subprocessors():
    from radar.export import VendorResult, providers_text
    from radar.contracts.models import VendorVerdict
    verdict = VendorVerdict(vendor_id="V-001", ai_usage="Yes", strength="Strong", coverage=1.0,
                            providers=["OpenAI", "GPT-4"])
    result = VendorResult(vendor_id="V-001", verdict=verdict)
    value = providers_text(result)
    assert "OpenAI" not in value and "GPT-4" not in value
    assert "subprocessor" in value.lower()


def test_explicit_subprocessor_in_evidence_is_populated():
    from radar.export import VendorResult, providers_text
    from radar.contracts.models import Signal, VendorVerdict
    signal = Signal(
        signal_id="S1", vendor_id="V-001", doc_id="D1", source_type="TRUST",
        url="https://x.test/trust",
        excerpt="OpenAI is authorized as a third-party subprocessor for customer inquiries.",
        char_start=0, char_end=74, category="chatbot",
        providers=["OpenAI"], specificity=0.8, weight=4.0, score=3.2,
        tag="STRONG", verified=True, method="RULE",
    )
    verdict = VendorVerdict(vendor_id="V-001", ai_usage="Yes", strength="Strong", coverage=1.0,
                            providers=["OpenAI"])
    result = VendorResult(vendor_id="V-001", verdict=verdict, signals=[signal])
    value = providers_text(result)
    assert "Explicitly verified AI subprocessor(s): OpenAI." == value


def test_general_ai_model_mention_is_not_subprocessor():
    from radar.export import VendorResult, providers_text
    from radar.contracts.models import Signal, VendorVerdict
    signal = Signal(
        signal_id="S1", vendor_id="V-001", doc_id="D1", source_type="WEB",
        url="https://x.test/product",
        excerpt="Our platform leverages GPT-4 and Llama algorithms for data analysis.",
        char_start=0, char_end=67, category="analysis",
        providers=["OpenAI"], specificity=0.6, weight=1.5, score=0.9,
        tag="PROBABLE", verified=True, method="RULE",
    )
    verdict = VendorVerdict(vendor_id="V-001", ai_usage="Yes", strength="Probable", coverage=1.0,
                            providers=["OpenAI"])
    result = VendorResult(vendor_id="V-001", verdict=verdict, signals=[signal])
    value = providers_text(result)
    assert "Explicitly verified AI subprocessor(s)" not in value
    assert "No named AI subprocessor explicitly verified in the collected evidence." in value


def test_local_llm_caches_only_valid_responses(tmp_path):
    llm = LocalLLM("http://127.0.0.1:11434", "m", tmp_path, session=FakeLLMSession(good_content("not present quote")))
    res = llm.propose(TECH)
    assert res is None
    assert llm.stats["invalid"] == 1
    # Cache directory should not have written the invalid proposal
    assert not list(tmp_path.glob("*.json"))

