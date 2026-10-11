from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from m4_intel.export_workbook import export_workbook


def _write_jsonl(path: Path, rows: list[dict]) -> None:
    import json
    path.write_text("".join(json.dumps(row) + "\n" for row in rows), encoding="utf-8")



def test_export_creates_separate_copy_and_summary_sheet(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendor Inventory"

    headers = [
        "Vendor ID",
        "Vendor Name",
        "Domain",
        "Service Description",
        "Business Criticality",
        "Data Sensitivity",
        "AI Usage Notes",
        "Assessment Status",
        "Contact",
        "Notes",
    ]

    for column, header in enumerate(headers, start=2):
        ws.cell(row=4, column=column, value=header)

    ws["B5"] = "V-000"
    ws["C5"] = "Example"
    ws["D5"] = "example.test"
    ws["B6"] = "V-001"
    ws["C6"] = "Vendor One"
    ws["D6"] = "vendor.test"

    wb.save(source)
    wb.close()

    manifest = tmp_path / "manifest.jsonl"
    candidates = tmp_path / "candidates.jsonl"

    _write_jsonl(manifest, [{
        "vendor_id": "V-001",
        "doc_id": "D1",
        "source_type": "WEB",
        "url": "https://vendor.test/",
        "text": "A sentence about AI.",
    }])

    _write_jsonl(candidates, [{
        "vendor_id": "V-001",
        "doc_id": "D1",
        "text": "A sentence about AI.",
        "char_start": 0,
        "char_end": 20,
        "categories": ["techniques"],
    }])

    output = tmp_path / "out.xlsx"
    result = export_workbook(source, manifest, candidates, output)

    assert result["vendors_exported"] == 1
    assert result["original_workbook_modified"] is False

    original = load_workbook(source, read_only=True)
    exported = load_workbook(output, read_only=True)

    assert "AI Radar Summary" not in original.sheetnames
    assert "AI Radar Summary" in exported.sheetnames

    rows = list(exported["AI Radar Summary"].iter_rows(values_only=True))
    print("SUMMARY SHEET ROWS:")
    for row in rows:
        print(repr(row))

    assert rows[1][0] == "V-001"
    assert rows[1][5] == "Inconclusive"
    assert all(row[0] != "V-000" for row in rows[1:])

    original.close()
    exported.close()
