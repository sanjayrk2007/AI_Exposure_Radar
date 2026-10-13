from openpyxl import Workbook
import pytest

from radar.m1_app.excel_io import read_vendor_profiles
from radar.m1_app.pipeline import run_vendor
from radar.contracts.models import VendorProfile


HEADERS = [
    "Vendor ID",
    "Vendor Name",
    "Vendor Description",
    "Service / Product Provided",
    "Category",
    "Public Website",
    "Business Process Supported",
    "Operational Dependency",
    "Data Classification Accessed",
    "Data Volume (annual)",
]

VENDOR_ROW = [
    "V-101",
    "Test Vendor",
    "A test vendor description",
    "Cloud software",
    "Technology",
    "https://example.com",
    "Internal reporting",
    "Moderate",
    "Internal data",
    "1000 records",
]


def make_workbook(path, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vendor Inventory"

    # Match the real workbook layout: headers in B4:K4.
    for column, header in enumerate(HEADERS, start=2):
        sheet.cell(row=4, column=column, value=header)

    for row_number, row in enumerate(rows, start=5):
        for column, value in enumerate(row, start=2):
            sheet.cell(row=row_number, column=column, value=value)

    workbook.save(path)
    workbook.close()


def test_reads_vendor_using_actual_workbook_headers(tmp_path):
    path = tmp_path / "vendors.xlsx"
    make_workbook(path, [VENDOR_ROW])

    vendors = read_vendor_profiles(path)

    assert len(vendors) == 1
    assert vendors[0].vendor_id == "V-101"
    assert vendors[0].name == "Test Vendor"
    assert vendors[0].description == "A test vendor description"
    assert vendors[0].website == "https://example.com"


def test_excludes_example_vendor(tmp_path):
    path = tmp_path / "vendors.xlsx"
    example = ["V-000", *VENDOR_ROW[1:]]
    make_workbook(path, [example, VENDOR_ROW])

    vendors = read_vendor_profiles(path)

    assert [vendor.vendor_id for vendor in vendors] == ["V-101"]


def test_rejects_unexpected_headers(tmp_path):
    path = tmp_path / "vendors.xlsx"
    make_workbook(path, [VENDOR_ROW])

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vendor Inventory"
    for column, header in enumerate(HEADERS, start=2):
        sheet.cell(row=4, column=column, value=header)
    sheet["D4"] = "Wrong Header"
    workbook.save(path)
    workbook.close()

    with pytest.raises(ValueError, match="Unexpected headers"):
        read_vendor_profiles(path)


def test_rejects_duplicate_vendor_ids(tmp_path):
    path = tmp_path / "vendors.xlsx"
    duplicate = VENDOR_ROW.copy()
    duplicate[1] = "Another Name"
    make_workbook(path, [VENDOR_ROW, duplicate])

    with pytest.raises(ValueError, match="Duplicate Vendor ID"):
        read_vendor_profiles(path)


def test_pipeline_rejects_example_vendor():
    vendor = VendorProfile(
        vendor_id="V-000",
        name="Example",
        description="Example vendor",
        service="Example service",
        category="Example",
        website="https://example.com",
        process="Example process",
        dependency="Low",
        data_class="Public",
        data_volume="Small",
    )

    with pytest.raises(ValueError, match="example row"):
        run_vendor(vendor)


def test_pipeline_rejects_invalid_mode():
    vendor = VendorProfile(
        vendor_id="V-101",
        name="Test Vendor",
        description="A test vendor",
        service="Cloud software",
        category="Technology",
        website="https://example.com",
        process="Internal reporting",
        dependency="Moderate",
        data_class="Internal data",
        data_volume="1000 records",
    )

    with pytest.raises(ValueError, match="mode"):
        run_vendor(vendor, mode="invalid")
